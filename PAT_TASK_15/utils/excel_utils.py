from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from PAT_TASK_15.utils.config_reader import ConfigReader

class ExcelUtils:

    @staticmethod
    def read_test_data():
        config_details = ConfigReader.get_config()
        excel_data_file = config_details["excel_data_file_location"]
        wb = load_workbook(excel_data_file)
        sheet = wb.active
        data = []

        for row_index in range(2, sheet.max_row +1):
            test_id = sheet.cell(row=row_index,column=1).value
            username = sheet.cell(row=row_index,column=2).value
            password = sheet.cell(row=row_index,column=3).value

            #If only there is a value in test_id, we are appending the data
            if test_id:
                data.append((row_index,test_id,username,password))

        wb.close()
        return data

    @staticmethod
    def update_test_results(row_index,status,date_str,time_str):
        config_details = ConfigReader.get_config()
        excel_data_file = config_details["excel_data_file_location"]
        wb = load_workbook(excel_data_file)
        sheet = wb.active

        #Color patterns for Pass/Fail Status
        green_fill = PatternFill(start_color='C6EFCE',end_color='C6EFCE',fill_type='solid')
        red_fill = PatternFill(start_color='FFC7CE',end_color='FFC7CE',fill_type='solid')

        #Updating Date (Column -4) and Time (Column - 5)
        sheet.cell(row=row_index,column=4,value=date_str)
        sheet.cell(row=row_index,column=5,value=time_str)

        #Updating the test result
        result_cell = sheet.cell(row=row_index,column=7,value=status)

        if status == 'Pass':
            result_cell.fill = green_fill
        else:
            result_cell.fill = red_fill

        wb.save(excel_data_file)
        wb.close()







